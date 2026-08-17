from collections import defaultdict
from decimal import Decimal, InvalidOperation
from django.shortcuts import get_object_or_404, render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_POST
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from django.db.models.functions import Substr, Cast
from django.db.models import IntegerField, Max
from decimal import Decimal, ROUND_HALF_UP
from .models import Producto, Categoria
from django.utils import timezone
from datetime import timedelta


def parse_decimal(value):
    """
    Convierte una cadena a Decimal, reemplazando comas por puntos y eliminando comillas raras.
    Si viene vacía o inválida, devuelve 0.
    """
    if not value:
        return Decimal('0')
    value = str(value).replace(',', '.').replace('“', '').replace('”', '').strip()
    try:
        return Decimal(value)
    except InvalidOperation:
        return Decimal('0')


@login_required
def lista_productos(request):
    productos = Producto.objects.filter(empresa=request.user.empresa)
    categorias = Categoria.objects.filter(empresa=request.user.empresa)
    
    # Nueva — para el modal del PDF
    categorias_pdf = (
        Producto.objects
        .filter(empresa=request.user.empresa)
        .values_list('categoria__nombre', flat=True)
        .distinct()
        .order_by('categoria__nombre')
    )

    return render(request, 'productos/lista_productos.html', {
        'productos': productos,
        'categorias': categorias,
        'categorias_pdf': list(categorias_pdf),  # nueva
    })
    
    
@login_required
def nuevo_producto(request):
    if request.method == 'POST' and request.headers.get('X-Requested-With', '').lower() == 'xmlhttprequest':
        nombre = request.POST.get('nombre')
        codigo = request.POST.get('codigo')
        categoria_id = request.POST.get('categoria')
        tipo_venta = request.POST.get('tipo_venta')

        precio_venta = parse_decimal(request.POST.get('precio_venta'))
        precio_compra = parse_decimal(request.POST.get('precio_compra', 0))
        stock_actual = parse_decimal(request.POST.get('stock_actual'))

        aplica_descuento = 'aplica_descuento' in request.POST
        if aplica_descuento:
            cantidad_minima_descuento = request.POST.get('cantidad_minima_descuento')
            precio_descuento_manual = parse_decimal(request.POST.get('precio_descuento_manual'))
        else:
            cantidad_minima_descuento = None
            precio_descuento_manual = None
            
        vende_por_bulto = 'vende_por_bulto' in request.POST
        if vende_por_bulto:
            unidades_por_bulto = request.POST.get('unidades_por_bulto')
            precio_por_bulto = parse_decimal(request.POST.get('precio_por_bulto'))
        else:
            unidades_por_bulto = None
            precio_por_bulto = None    

        alerta_stock_bajo = 'alerta_stock_bajo' in request.POST
        if alerta_stock_bajo:
            stock_minimo_alerta = request.POST.get('stock_minimo_alerta')
            stock_minimo_alerta = int(stock_minimo_alerta) if stock_minimo_alerta else 0
        else:
            stock_minimo_alerta = None

        fecha_vencimiento = request.POST.get('fecha_vencimiento') or None
        dias_aviso_str = request.POST.get('dias_aviso_vencimiento')
        dias_aviso_vencimiento = int(dias_aviso_str) if dias_aviso_str else 7

        if Producto.objects.filter(codigo=codigo, empresa=request.user.empresa).exists():
            return JsonResponse({'success': False, 'message': 'El código de producto ya existe para esta empresa.'})

        categoria = get_object_or_404(Categoria, id=categoria_id, empresa=request.user.empresa)

        producto = Producto.objects.create(
            nombre=nombre,
            codigo=codigo,
            categoria=categoria,
            precio_venta=precio_venta,
            precio_compra=precio_compra,
            stock_actual=stock_actual,
            empresa=request.user.empresa,
            tipo_venta=tipo_venta,
            aplica_descuento=aplica_descuento,
            cantidad_minima_descuento=int(cantidad_minima_descuento) if cantidad_minima_descuento else 0,
            precio_descuento_manual=precio_descuento_manual,
            vende_por_bulto=vende_por_bulto,
            unidades_por_bulto=int(unidades_por_bulto) if unidades_por_bulto else None,
            precio_por_bulto=precio_por_bulto,
            alerta_stock_bajo=alerta_stock_bajo,
            stock_minimo_alerta=stock_minimo_alerta,
            fecha_vencimiento=fecha_vencimiento,
            dias_aviso_vencimiento=dias_aviso_vencimiento,
        )

        return JsonResponse({
            'success': True,
            'producto': {
                'id': producto.id,
                'nombre': producto.nombre,
                'codigo': producto.codigo,
                'categoria': producto.categoria.nombre,
                'categoria_id': producto.categoria.id,
                'precio_venta': float(producto.precio_venta),
                'stock_actual': float(producto.stock_actual),
                'tipo_venta': producto.tipo_venta,
                'alerta_stock_bajo': producto.alerta_stock_bajo,
                'stock_minimo_alerta': producto.stock_minimo_alerta,
                'fecha_vencimiento': producto.fecha_vencimiento.isoformat() if producto.fecha_vencimiento else None,
                'vende_por_bulto': producto.vende_por_bulto,                                   
                'unidades_por_bulto': producto.unidades_por_bulto,                            
                'precio_por_bulto': float(producto.precio_por_bulto) if producto.precio_por_bulto else None,
            }
        })

    return JsonResponse({'success': False, 'message': 'Método no permitido o no es una solicitud AJAX.'})


from decimal import Decimal, InvalidOperation
from django.shortcuts import get_object_or_404, render, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required

from collections import defaultdict
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
 

# Si ya tenés parse_decimal y querés seguir usándolo, puedes usarlo en lugar
# de Decimal(...) después de limpiar la cadena. Aquí uso Decimal directamente.

def _clean_number_string(s):
    """
    Limpia una cadena numérica en formato human-readable:
    - "1.700" -> "1700"
    - "12.450,75" -> "12450.75"
    - "4,5" -> "4.5"
    Devuelve string listo para Decimal(...) o None si s es vacío.
    """
    if s is None:
        return None
    s = str(s).strip()
    if s == '':
        return None
    # Eliminar espacios
    s = s.replace(' ', '')
    # Quitar puntos (separador de miles)
    s = s.replace('.', '')
    # Reemplazar coma decimal por punto
    s = s.replace(',', '.')
    return s


@login_required
def editar_producto(request, id):
    producto = get_object_or_404(
        Producto,
        id=id,
        empresa=request.user.empresa
    )

    if request.method == 'POST':
        producto.nombre = request.POST.get('nombre', producto.nombre)

        # --- VALIDACIÓN DE CÓDIGO DUPLICADO ---
        codigo_post = request.POST.get('codigo', '').strip()

        if Producto.objects.filter(
            empresa=request.user.empresa,
            codigo=codigo_post
        ).exclude(id=producto.id).exists():

            messages.error(request, "Ya existe un producto con ese código.")
            categorias = Categoria.objects.filter(empresa=request.user.empresa)
            return render(request, 'productos/editar_producto.html', {
                'producto': producto,
                'categorias': categorias,
            })

        # ✅ ASIGNAMOS EL CÓDIGO SOLO SI PASÓ LA VALIDACIÓN
        producto.codigo = codigo_post

        # --- PRECIOS Y STOCK ---
        raw_precio_venta = request.POST.get('precio_venta')
        raw_precio_compra = request.POST.get('precio_compra')
        raw_stock = request.POST.get('stock_actual')

        precio_venta_str = _clean_number_string(raw_precio_venta)
        precio_compra_str = _clean_number_string(raw_precio_compra)
        stock_str = _clean_number_string(raw_stock)

        try:
            if precio_venta_str is not None:
                producto.precio_venta = Decimal(precio_venta_str)

            if precio_compra_str is not None:
                producto.precio_compra = Decimal(precio_compra_str)

            if stock_str is not None:
                stock_nuevo = Decimal(stock_str)
                producto.stock_actual = max(stock_nuevo, Decimal('0'))

        except (InvalidOperation, ValueError):
            messages.error(
                request,
                'Formato de número inválido. Revisa precios y stock.'
            )
            categorias = Categoria.objects.filter(empresa=request.user.empresa)
            return render(request, 'productos/editar_producto.html', {
                'producto': producto,
                'categorias': categorias,
            })

        categoria_id = request.POST.get('categoria')
        if categoria_id:
            producto.categoria = Categoria.objects.get(
                id=categoria_id,
                empresa=request.user.empresa
            )

        producto.aplica_descuento = 'aplica_descuento' in request.POST

        if producto.aplica_descuento:
            raw_cant_min = request.POST.get('cantidad_minima_descuento')
            cant_min_str = _clean_number_string(raw_cant_min)

            try:
                producto.cantidad_minima_descuento = (
                    int(Decimal(cant_min_str))
                    if cant_min_str is not None else 0
                )
            except (InvalidOperation, ValueError, TypeError):
                producto.cantidad_minima_descuento = 0

            raw_precio_manual = request.POST.get('precio_descuento_manual')
            precio_manual_str = _clean_number_string(raw_precio_manual)

            try:
                producto.precio_descuento_manual = (
                    Decimal(precio_manual_str) if precio_manual_str is not None else None
                )
            except (InvalidOperation, ValueError, TypeError):
                producto.precio_descuento_manual = None
        else:
            producto.cantidad_minima_descuento = 0
            producto.precio_descuento_manual = None
        
        producto.vende_por_bulto = 'vende_por_bulto' in request.POST

        if producto.vende_por_bulto:
            raw_unidades_bulto = request.POST.get('unidades_por_bulto')
            unidades_bulto_str = _clean_number_string(raw_unidades_bulto)

            try:
                producto.unidades_por_bulto = (
                    int(Decimal(unidades_bulto_str))
                    if unidades_bulto_str is not None else None
                )
            except (InvalidOperation, ValueError, TypeError):
                producto.unidades_por_bulto = None

            raw_precio_bulto = request.POST.get('precio_por_bulto')
            precio_bulto_str = _clean_number_string(raw_precio_bulto)

            try:
                producto.precio_por_bulto = (
                    Decimal(precio_bulto_str)
                    if precio_bulto_str is not None else None
                )
            except (InvalidOperation, ValueError, TypeError):
                producto.precio_por_bulto = None
        else:
            producto.unidades_por_bulto = None
            producto.precio_por_bulto = None
        
        # --- ALERTA STOCK BAJO ---
        producto.alerta_stock_bajo = 'alerta_stock_bajo' in request.POST
        if producto.alerta_stock_bajo:
            raw_min_alerta = request.POST.get('stock_minimo_alerta')
            min_alerta_str = _clean_number_string(raw_min_alerta)
            try:
                producto.stock_minimo_alerta = (
                    int(Decimal(min_alerta_str)) if min_alerta_str else 0
                )
            except (InvalidOperation, ValueError, TypeError):
                producto.stock_minimo_alerta = 0
        else:
            producto.stock_minimo_alerta = None

        # --- FECHA DE VENCIMIENTO ---
        producto.fecha_vencimiento = request.POST.get('fecha_vencimiento') or None
        dias_aviso_str = request.POST.get('dias_aviso_vencimiento')
        producto.dias_aviso_vencimiento = int(dias_aviso_str) if dias_aviso_str else 7

        producto.save()
        messages.success(request, 'Producto actualizado correctamente.')
        return redirect('lista_productos')

    categorias = Categoria.objects.filter(empresa=request.user.empresa)
    return render(request, 'productos/editar_producto.html', {
        'producto': producto,
        'categorias': categorias,
    })


@login_required
def eliminar_producto(request, id):
    producto = get_object_or_404(Producto, id=id, empresa=request.user.empresa)
    producto.delete()
    messages.success(request, 'Producto eliminado correctamente.')
    return redirect('lista_productos')


# -------------------------------------------------------------------
# PASO 1 — Pantalla para seleccionar y ordenar categorías
# -------------------------------------------------------------------
@login_required
def seleccionar_categorias_pdf(request):
    """
    Muestra un formulario con las categorías disponibles.
    El usuario puede activar/desactivar y arrastrar para reordenar.
    Al confirmar, hace POST a exportar_productos_pdf.
    """
    from .models import Producto  # ajustá el import según tu proyecto

    categorias = (
        Producto.objects
        .filter(empresa=request.user.empresa)
        .values_list('categoria__nombre', flat=True)
        .distinct()
        .order_by('categoria__nombre')
    )

    return render(request, 'seleccionar_categorias_pdf.html', {
        'categorias': list(categorias),
    })


# -------------------------------------------------------------------
# PASO 2 — Generar el PDF con el orden y selección recibidos
# -------------------------------------------------------------------
@login_required
def exportar_productos_pdf(request):
    from .models import Producto  # ajustá el import según tu proyecto

    def formatear_precio(valor):
        return "${:,}".format(int(valor)).replace(",", ".")

    # ------------------------------------------------------------------
    # Leer categorías seleccionadas y su orden desde el POST
    # Si viene por GET (acceso directo), usar todas en orden alfabético
    # ------------------------------------------------------------------
    if request.method == 'POST':
        orden_raw = request.POST.get('categorias_orden', '')
        categorias_ordenadas = [c.strip() for c in orden_raw.split(',') if c.strip()]
    else:
        categorias_ordenadas = []

    productos_qs = Producto.objects.filter(
        empresa=request.user.empresa
    ).select_related('categoria')

    productos_por_categoria = defaultdict(list)
    for prod in productos_qs:
        productos_por_categoria[prod.categoria.nombre].append(prod)

    if not categorias_ordenadas:
        categorias_ordenadas = sorted(productos_por_categoria.keys())

    categorias_finales = [
        c for c in categorias_ordenadas
        if c in productos_por_categoria
    ]

    # ------------------------------------------------------------------
    # Generar PDF
    # ------------------------------------------------------------------
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="productos.pdf"'

    p = canvas.Canvas(response, pagesize=A4)
    width, height = A4
    y = height - 50

    # Columnas (right-aligned para los precios, con más espacio horizontal)
    X_CODIGO = 50
    X_NOMBRE = 135
    X_PRECIO = 370
    X_DESCUENTO = 460
    X_BULTO = 545  # cerca del margen derecho (width - 50 ≈ 545)

    def imprimir_titulo():
        nonlocal y
        p.setFont("Helvetica-Bold", 16)
        p.drawString(50, y, f"Lista de Productos - {request.user.empresa.nombre}")
        y -= 40

    def imprimir_encabezados_columnas():
        nonlocal y
        p.setFont("Helvetica-Bold", 12)
        p.drawString(X_CODIGO, y, "Código")
        p.drawString(X_NOMBRE, y, "Nombre")
        p.drawRightString(X_PRECIO, y, "Precio")
        p.drawRightString(X_DESCUENTO, y, "P. Desc.")
        p.drawRightString(X_BULTO, y, "P. Bulto")
        y -= 20

    imprimir_titulo()

    for categoria in categorias_finales:
        productos_categoria = productos_por_categoria[categoria]

        if y < 100:
            p.showPage()
            y = height - 50
            imprimir_titulo()

        p.setFont("Helvetica-Bold", 14)
        p.drawString(50, y, f"  {categoria}")
        y -= 25

        imprimir_encabezados_columnas()

        p.setFont("Helvetica", 10)

        for prod in productos_categoria:
            if y < 60:
                p.showPage()
                y = height - 50

                p.setFont("Helvetica-Bold", 14)
                p.drawString(50, y, f"  {categoria} (continuación)")
                y -= 25

                imprimir_encabezados_columnas()

                p.setFont("Helvetica", 10)

            p.drawString(X_CODIGO, y, str(prod.codigo))
            p.drawString(X_NOMBRE, y, prod.nombre)
            p.drawRightString(X_PRECIO, y, formatear_precio(prod.precio_venta))

            # Precio con descuento por cantidad (usa precio_descuento_manual, ver nota arriba)
            if prod.aplica_descuento and prod.precio_descuento_manual:
                texto_desc = formatear_precio(prod.precio_descuento_manual)
                if prod.cantidad_minima_descuento:
                    texto_desc += f" ({prod.cantidad_minima_descuento}+)"
            else:
                texto_desc = "-"
            p.drawRightString(X_DESCUENTO, y, texto_desc)

            # Precio por bulto
            if prod.vende_por_bulto and prod.precio_por_bulto:
                texto_bulto = formatear_precio(prod.precio_por_bulto)
                if prod.unidades_por_bulto:
                    texto_bulto += f" (x{prod.unidades_por_bulto})"
            else:
                texto_bulto = "-"
            p.drawRightString(X_BULTO, y, texto_bulto)

            p.line(50, y - 2, width - 50, y - 2)
            y -= 18

        y -= 15

    p.showPage()
    p.save()
    return response
@login_required
@require_POST
def nueva_categoria(request):
    nombre = request.POST.get('nombre')
    empresa = request.user.empresa
    if Categoria.objects.filter(nombre__iexact=nombre, empresa=empresa).exists():
        return JsonResponse({'success': False, 'message': 'La categoría ya existe.'})

    cat = Categoria.objects.create(nombre=nombre, empresa=empresa)
    return JsonResponse({'success': True, 'id': cat.id, 'nombre': cat.nombre})


# ESTO VA EN TUS VISTAS (views.py)
# Actualiza tu función crear_categoria existente O agrega esta si no la tienes en la forma correcta

import json
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods

@login_required
@require_http_methods(["POST"])
def crear_categoria(request):
    """
    Endpoint AJAX para crear una nueva categoría desde editar_producto.html
    
    Acepta:
    - application/x-www-form-urlencoded (POST tradicional)
    - application/json (AJAX con JSON)
    
    Retorna JSON: { "success": bool, "categoria": {...}, "error": string }
    """
    
    try:
        # Intentar parsear como JSON (desde AJAX)
        if request.content_type == 'application/json':
            data = json.loads(request.body)
            nombre = data.get('nombre', '').strip()
        else:
            # Si es POST tradicional
            nombre = request.POST.get('nombre', '').strip()

        # Validación: nombre requerido
        if not nombre:
            return JsonResponse({
                'success': False,
                'error': 'El nombre de la categoría es requerido'
            })

        # Validación: nombre no demasiado largo
        if len(nombre) > 100:
            return JsonResponse({
                'success': False,
                'error': 'El nombre es demasiado largo (máximo 100 caracteres)'
            })

        # Validación: no existe categoría con el mismo nombre (case-insensitive)
        if Categoria.objects.filter(
            empresa=request.user.empresa,
            nombre__iexact=nombre
        ).exists():
            return JsonResponse({
                'success': False,
                'error': f'La categoría "{nombre}" ya existe'
            })

        # Crear la nueva categoría
        nueva_categoria = Categoria.objects.create(
            nombre=nombre,
            empresa=request.user.empresa
        )

        return JsonResponse({
            'success': True,
            'categoria': {
                'id': nueva_categoria.id,
                'nombre': nueva_categoria.nombre
            }
        })

    except json.JSONDecodeError:
        return JsonResponse({
            'success': False,
            'error': 'Formato de datos inválido'
        }, status=400)
    
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f'Error al crear la categoría: {str(e)}'
        }, status=500)

@login_required
def buscar_producto_por_codigo(request):
    codigo = request.GET.get('codigo')
    empresa = request.user.empresa
    try:
        producto = Producto.objects.get(codigo=codigo, empresa=empresa)
        return JsonResponse({
            'success': True,
            'nombre': producto.nombre,
            'precio_venta': float(producto.precio_venta),
            'stock_actual': float(producto.stock_actual),
            'tipo_venta': producto.tipo_venta,
            'codigo': producto.codigo,
        })
    except Producto.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Producto no encontrado'})



@login_required
def generar_codigo_producto(request):
    empresa = request.user.empresa
    producto_id = request.GET.get("producto_id")

    qs = Producto.objects.filter(
        empresa=empresa,
        codigo__startswith="INT"
    )

    if producto_id:
        qs = qs.exclude(id=producto_id)

    # ⬇️ EXTRAER PARTE NUMÉRICA DEL CÓDIGO
    qs = qs.annotate(
        numero_codigo=Cast(Substr("codigo", 4), IntegerField())
    )

    max_numero = qs.aggregate(max_num=Max("numero_codigo"))["max_num"]

    siguiente = (max_numero or 0) + 1
    codigo = f"INT{siguiente:06d}"

    return JsonResponse({
        "success": True,
        "codigo": codigo
    })



@login_required
@require_POST
def actualizar_producto_inline(request):
    """
    Actualiza un producto inline vía AJAX
    """
    try:
        producto_id = request.POST.get('producto_id')
        campo = request.POST.get('campo')  # 'precio_compra', 'precio_venta', 'stock_actual'
        valor = request.POST.get('valor')
        
        producto = get_object_or_404(
            Producto, 
            id=producto_id, 
            empresa=request.user.empresa
        )
        
        # Limpiar y convertir valor
        valor_limpio = _clean_number_string(valor)
        
        if valor_limpio is None:
            return JsonResponse({
                'success': False, 
                'message': 'Valor inválido'
            })
        
        # Actualizar según el campo
        if campo == 'precio_compra':
            producto.precio_compra = Decimal(valor_limpio)
        elif campo == 'precio_venta':
            producto.precio_venta = Decimal(valor_limpio)
        elif campo == 'stock_actual':
            producto.stock_actual = max(Decimal(valor_limpio), Decimal('0'))
        else:
            return JsonResponse({
                'success': False, 
                'message': 'Campo no válido'
            })
        
        producto.save()
        
        # Formatear valor para respuesta
        valor_formateado = "{:,.2f}".format(float(getattr(producto, campo))).replace(",", "X").replace(".", ",").replace("X", ".")
        
        return JsonResponse({
            'success': True,
            'valor_formateado': valor_formateado,
            'message': 'Actualizado correctamente'
        })
        
    except (InvalidOperation, ValueError) as e:
        return JsonResponse({
            'success': False, 
            'message': 'Formato de número inválido'
        })
    except Exception as e:
        return JsonResponse({
            'success': False, 
            'message': str(e)
        })
        

@login_required
def lista_productos_edicion_masiva(request):
    from .models import Categoria
    productos = Producto.objects.filter(empresa=request.user.empresa)
    categorias = Categoria.objects.filter(empresa=request.user.empresa)
    return render(request, 'productos/lista_productos_edicion_masiva.html', {
        'productos': productos,
        'categorias': categorias,
    })     


@login_required
@require_POST
def actualizar_precios_masivo(request):
    """
    Actualiza precios de múltiples productos con opciones de redondeo
    """
    try:
        ids = request.POST.getlist('ids[]')
        porcentaje = Decimal(request.POST.get('porcentaje'))
        campo = request.POST.get('campo')
        redondear = request.POST.get('redondear') == '1'
        redondear_centenas = request.POST.get('redondear_centenas') == '1'
        margen_str = request.POST.get('margen_ganancia', '').strip()
 
        productos = Producto.objects.filter(id__in=ids, empresa=request.user.empresa)
        factor = 1 + porcentaje / 100
 
        def aplicar_redondeo(valor):
            """
            Aplica redondeo según las opciones seleccionadas
            - redondear: redondea a entero (1, 2, 3, 4, 5...)
            - redondear_centenas: redondea a centenas (100, 200, 300, 400, 500...)
            """
            valor_float = float(valor)
            
            if redondear_centenas:
                # Redondea a la centena más cercana
                # 653 -> 700, 643 -> 600, 550 -> 600, 549 -> 500
                return Decimal(str(round(valor_float / 100) * 100))
            elif redondear:
                # Redondea a entero
                return valor.quantize(Decimal('1'), rounding=ROUND_HALF_UP)
            else:
                # Mantiene 2 decimales
                return valor.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
 
        for p in productos:
            if campo == 'ambos' and margen_str:
                # Sube precio compra y recalcula venta con el margen indicado
                margen = Decimal(margen_str) / 100
                nuevo_compra = p.precio_compra * factor
                p.precio_compra = aplicar_redondeo(nuevo_compra)
                
                nuevo_venta = nuevo_compra / (1 - margen)
                p.precio_venta = aplicar_redondeo(nuevo_venta)
 
            else:
                if campo in ('precio_compra', 'ambos'):
                    p.precio_compra = aplicar_redondeo(p.precio_compra * factor)
                if campo in ('precio_venta', 'ambos'):
                    p.precio_venta = aplicar_redondeo(p.precio_venta * factor)
 
            p.save()
 
        return JsonResponse({
            'success': True,
            'actualizados': productos.count(),
            'mensaje': f'{productos.count()} producto(s) actualizado(s) correctamente'
        })
 
    except ValueError as e:
        return JsonResponse({
            'success': False,
            'message': f'Valor inválido: {str(e)}'
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': str(e)
        })



from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

@login_required
def exportar_productos_excel(request):
    """
    Exporta la lista de productos a un archivo Excel (.xlsx)
    agrupados por categoría, con formato profesional.
    """
    
    # Obtener productos ordenados por categoría
    productos = Producto.objects.filter(
        empresa=request.user.empresa
    ).order_by('categoria__nombre', 'nombre')
    
    # Agrupar por categoría
    productos_por_categoria = defaultdict(list)
    for prod in productos:
        productos_por_categoria[prod.categoria.nombre].append(prod)
    
    # Crear libro de Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    
    # Estilos
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    categoria_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    categoria_font = Font(bold=True, size=13)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título principal
    ws.merge_cells('A1:G1')
    ws['A1'] = f"📄 Lista de Productos - {request.user.empresa.nombre}"
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    fila = 3  # Empezar después del título
    
    # Recorrer categorías
    for categoria, productos_categoria in productos_por_categoria.items():
        
        # Nombre de la categoría
        ws.merge_cells(f'A{fila}:G{fila}')
        celda_categoria = ws[f'A{fila}']
        celda_categoria.value = f"▶ {categoria}"
        celda_categoria.font = categoria_font
        celda_categoria.fill = categoria_fill
        celda_categoria.alignment = Alignment(horizontal='left', vertical='center')
        fila += 1
        
        # Encabezados de tabla
        encabezados = ['Código', 'Nombre', 'Stock', 'Tipo', 'Precio', '% Extra envío', 'Precio Final']
        for col_num, encabezado in enumerate(encabezados, 1):
            celda = ws.cell(row=fila, column=col_num)
            celda.value = encabezado
            celda.font = header_font
            celda.fill = header_fill
            celda.alignment = Alignment(horizontal='center', vertical='center')
            celda.border = border
        
        fila += 1
        
        # Productos de la categoría
        for prod in productos_categoria:
            # Formatear stock según tipo
            if prod.tipo_venta == "unidad":
                stock_valor = int(prod.stock_actual)
            else:
                stock_valor = round(prod.stock_actual, 2)
            
            # Escribir datos
            ws.cell(row=fila, column=1, value=prod.codigo).border = border
            ws.cell(row=fila, column=2, value=prod.nombre).border = border
            
            # Stock
            celda_stock = ws.cell(row=fila, column=3, value=stock_valor)
            celda_stock.border = border
            celda_stock.alignment = Alignment(horizontal='center')
            if prod.tipo_venta != "unidad":
                celda_stock.number_format = '0.00'
            
            # Tipo de venta
            celda_tipo = ws.cell(row=fila, column=4, value=prod.get_tipo_venta_display())
            celda_tipo.border = border
            celda_tipo.alignment = Alignment(horizontal='center')
            
            # Precio
            celda_precio = ws.cell(row=fila, column=5, value=int(prod.precio_venta))
            celda_precio.border = border
            celda_precio.alignment = Alignment(horizontal='right')
            celda_precio.number_format = '"$"#,##0'
            
            # COLUMNA F: % Aumento (editable, vacía por defecto)
            celda_porcentaje = ws.cell(row=fila, column=6)
            celda_porcentaje.border = border
            celda_porcentaje.alignment = Alignment(horizontal='center')
            celda_porcentaje.number_format = '0"%"'  # Muestra "5%" cuando escribe 5
            
            # COLUMNA G: Precio Final (fórmula automática)
            celda_final = ws.cell(row=fila, column=7)
            celda_final.border = border
            celda_final.alignment = Alignment(horizontal='right')
            celda_final.number_format = '"$"#,##0'
            # Fórmula: si F está vacío muestra E, si no calcula E + (E * F/100)
            celda_final.value = f'=IF(F{fila}="",E{fila},E{fila}+(E{fila}*F{fila}/100))'
            
            fila += 1
        
        fila += 1  # Espacio entre categorías
    
    # Ajustar anchos de columna
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 15
    
    # Preparar respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="productos.xlsx"'
    
    wb.save(response)
    return response